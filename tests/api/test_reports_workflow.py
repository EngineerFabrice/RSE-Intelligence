import io

from openpyxl import load_workbook

from tests.fixtures.sample_report import build_sample_pdf

from app.extensions import db
from app.ingestion.pipeline import process_report
from app.models.report import Report
from app.services import duplicate_detection


def _seed_processed_report(app, tmp_path, mismatch=False):
    pdf_path = tmp_path / "sample.pdf"
    build_sample_pdf(pdf_path, shares_traded_mismatch=mismatch)
    file_bytes = pdf_path.read_bytes()

    report = Report(filename="sample.pdf", original_filename="sample.pdf",
                     file_hash=duplicate_detection.compute_file_hash(file_bytes), storage_path=str(pdf_path))
    db.session.add(report)
    db.session.commit()
    process_report(report.id)
    db.session.refresh(report)
    return report


def test_list_reports_endpoint(admin_client, app, tmp_path):
    _seed_processed_report(app, tmp_path)
    resp = admin_client.get("/api/reports")
    assert resp.status_code == 200
    assert resp.get_json()["data"]["total"] >= 1


def test_preview_endpoint_returns_all_sections(admin_client, app, tmp_path):
    report = _seed_processed_report(app, tmp_path)
    resp = admin_client.get(f"/api/reports/{report.id}/preview")
    body = resp.get_json()["data"]
    assert len(body["stock"]) == 2
    assert body["market_stats"] is not None
    assert len(body["bonds"]) == 1


def test_approve_blocked_when_critical_issues_open(admin_client, app, tmp_path):
    report = _seed_processed_report(app, tmp_path, mismatch=True)
    resp = admin_client.post(f"/api/reports/{report.id}/approve")
    assert resp.status_code == 409


def test_resolve_issue_then_approve_succeeds(admin_client, app, tmp_path):
    report = _seed_processed_report(app, tmp_path, mismatch=True)

    issues_resp = admin_client.get(f"/api/reports/{report.id}/preview")
    issues = issues_resp.get_json()["data"]["issues"]
    open_critical = [i for i in issues if i["severity"] == "critical" and not i["resolution"]]
    for issue in open_critical:
        r = admin_client.post(f"/api/reports/{report.id}/issues/{issue['id']}/resolve",
                               json={"resolution": "chosen_a"})
        assert r.status_code == 200

    resp = admin_client.post(f"/api/reports/{report.id}/approve")
    assert resp.status_code == 200
    assert resp.get_json()["data"]["approval_status"] == "approved"


def test_export_excel_after_approval(admin_client, app, tmp_path):
    report = _seed_processed_report(app, tmp_path, mismatch=False)
    admin_client.post(f"/api/reports/{report.id}/approve")

    resp = admin_client.get(f"/api/reports/{report.id}/export/excel")
    assert resp.status_code == 200
    assert resp.data[:4] == b"PK\x03\x04"


def test_preview_and_excel_share_canonical_equity_values(
        admin_client, app, tmp_path):
    report = _seed_processed_report(app, tmp_path, mismatch=False)
    preview = admin_client.get(
        f"/api/reports/{report.id}/preview"
    ).get_json()["data"]
    workbook_bytes = admin_client.get(
        f"/api/reports/{report.id}/export/excel"
    ).data
    sheet = load_workbook(io.BytesIO(workbook_bytes), data_only=True)["STOCK"]

    preview_by_symbol = {row["symbol"]: row for row in preview["stock"]}
    excel_by_symbol = {
        row[0]: row for row in sheet.iter_rows(min_row=2, values_only=True)
    }
    assert set(preview_by_symbol) == set(excel_by_symbol)
    for symbol, preview_row in preview_by_symbol.items():
        excel_row = excel_by_symbol[symbol]
        assert float(preview_row["closing_price"]) == excel_row[1]
        assert preview_row["volume"] == excel_row[2]
        assert float(preview_row["value_turnover"]) == excel_row[3]


def test_market_overview_reflects_approved_report(admin_client, app, tmp_path):
    report = _seed_processed_report(app, tmp_path, mismatch=False)
    admin_client.post(f"/api/reports/{report.id}/approve")

    resp = admin_client.get("/api/market/overview")
    body = resp.get_json()["data"]
    assert body["report"]["id"] == report.id
    assert "narrative" in body


def test_report_endpoint_exposes_progress_for_the_upload_ui(admin_client, app, tmp_path):
    # Finished processing (whether or not it needs review) must report 100%/"done" —
    # the upload page's progress bar only shows "Report Ready" once this is true.
    report = _seed_processed_report(app, tmp_path, mismatch=False)
    resp = admin_client.get(f"/api/reports/{report.id}")
    progress = resp.get_json()["data"]["progress"]
    assert progress["percent"] == 100
    assert progress["state"] == "done"
    # No internal step vocabulary leaks into the payload consumed by the upload page.
    assert "Equities identified" not in str(progress)


def test_home_is_only_the_upload_workflow(admin_client):
    response = admin_client.get("/")
    assert response.status_code == 200
    page = response.get_data(as_text=True)
    assert "Upload RSE Market Report" in page
    assert "Market Dashboard" not in page
    assert "turnoverChart" not in page
    assert "Confidence" not in page


def test_reports_browse_is_read_only_and_verified_only(admin_client, app, tmp_path):
    report = _seed_processed_report(app, tmp_path)
    response = admin_client.get("/reports/browse")
    assert response.status_code == 200
    assert report.name not in response.get_data(as_text=True)

    admin_client.post(f"/api/reports/{report.id}/approve")
    response = admin_client.get("/reports/browse")
    page = response.get_data(as_text=True)
    assert response.status_code == 200
    assert report.name in page
    assert "Reprocess" not in page
    assert "Delete" not in page


def test_preview_page_uses_report_sections_without_confidence_ui(admin_client, app, tmp_path):
    report = _seed_processed_report(app, tmp_path)
    response = admin_client.get(f"/reports/{report.id}/preview")
    page = response.get_data(as_text=True)
    assert response.status_code == 200
    assert "Equities" in page
    assert "Indices" in page
    assert "Market Stats" in page
    assert "Confidence" not in page
